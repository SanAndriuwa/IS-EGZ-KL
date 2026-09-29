"""Insert the teacher-feedback addendum into the existing formatted DOCX.

The full source remains build_colloquium_pdf.py. This small updater preserves
the already rendered equations when a local TeX compiler is unavailable.
"""
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
path = ROOT / 'docs' / 'KOLIOKVIUMO_PLANAS.docx'
doc = Document(path)
title = '3.6. Naujausi literatūroje taikomi metodai — atnaujinimas pagal dėstytojo pastabas'
if any(p.text == title for p in doc.paragraphs):
    paragraph = next(p for p in doc.paragraphs if p.text.startswith('Papildomų bandymų planas:'))
    paragraph.paragraph_format.space_before = Pt(10)
    doc.save(path)
    raise SystemExit('Updated addendum spacing.')
anchor = next(p for p in doc.paragraphs if p.text.startswith('4. Pagrindinis sprendimas'))


def insert(text, style='Normal'):
    paragraph = anchor.insert_paragraph_before(text, style=style)
    return paragraph


insert(title, 'Heading 2')
insert('Šis poskyris pridėtas po pirminio plano; jis nekeičia anksčiau pasirinktų RF, LR, GB ir H1 statuso. Abdullah-All-Tanvir ir kt. (2023) artimam sesijų pirkimo ketinimo uždaviniui taiko XGBoost su požymių atranka ir balansavimu. Satu ir Islam (2023) tam pačiam UCI rinkiniui tiria RF, transformacijas, SMOTE ir požymių atranką. Setyawan ir Himawan (2026) lygina XGBoost, LightGBM ir CatBoost su SMOTE bei platesne optimizacija. Tai artimi uždaviniai, bet jų vertinimo taisyklės ir galimai kitokie požymiai neįrodo geresnio AP mūsų laikiniame skaidyme.')

rows = [
    ['Svarstytas kelias', 'Kodėl tinka', 'Sprendimas šiame plane'],
    ['CPU XGBoost', 'Tinka sesijų lentelei; klasės svoris tikrina disbalanso įtaką.', 'Pasirinktas vienam ribotam bandymui; telpa į CPU biudžetą, bet brangesnis už LR.'],
    ['RF su SMOTE ir atranka', 'Galėtų mažinti disbalanso ir nereikalingų požymių poveikį.', 'Dabar atmestas: kelios keičiamos grandys, sintetinės kategorijos ir sunkesnė interpretacija.'],
    ['Stacking / balsavimas', 'Gali jungti skirtingų modelių signalus.', 'Atmestas: papildomas meta-modelio skaidymas, kaina ir sudėtingesnis gynimas.'],
]
table = doc.add_table(rows=len(rows), cols=3)
table.style = 'Table Grid'
for i, row in enumerate(rows):
    table.rows[i]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    for cell, value in zip(table.rows[i].cells, row):
        cell.text = value
        for run in cell.paragraphs[0].runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(9)
            run.bold = i == 0
table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
anchor._p.addprevious(table._tbl)
closing = insert('Papildomų bandymų planas: nekeisti pagrindinio eksperimento. Po jo atskirai mokyti XGBoost ir RF parametrų variantus tik train; kiekvienos šeimos kandidatą parinkti pagal validation AP, slenkstį — pagal validation F2, testą vertinti tik po pasirinkimo. Išsaugoti visų kandidatų validacijos AP ir neigiamą rezultatą, jei pagerėjimo nėra. Kadangi pagrindinis testas jau žinomas, tai būtų post-test exploratory bandymas; nepriklausomam patvirtinimui reikėtų naujo būsimo laikotarpio. Breiman (2001) ir Friedman (2001) paaiškina mechanizmą, bet nėra šiuolaikinio pranašumo šiam rinkiniui įrodymas. Detalesnė šaltinių patikra pateikta NAUJA_LITERATURA.md.')
closing.paragraph_format.space_before = Pt(10)

reference_anchor = next(p for p in doc.paragraphs if p.text.startswith('Breiman, L. (2001)'))
references = [
    'Abdullah-All-Tanvir, Khandokar, I. A., Islam, A. K. M. M., Islam, S., & Shatabda, S. (2023). A gradient boosting classifier for purchase intention prediction of online shoppers. Heliyon, 9(4), e15163. https://doi.org/10.1016/j.heliyon.2023.e15163',
    'Satu, M. S., & Islam, S. F. (2023). Modeling online customer purchase intention behavior applying different feature engineering and classification techniques. Discover Artificial Intelligence, 3, 36. https://doi.org/10.1007/s44163-023-00086-0',
    'Setyawan, I. B., & Himawan, H. (2026). Optimasi Bayesian pada Gradient Boosting untuk Prediksi Niat Beli E-Commerce pada Dataset dengan Ketidakseimbangan Kelas. Building of Informatics, Technology and Science, 8(1), 51–61. https://doi.org/10.47065/bits.v8i1.9710',
]
for text in references:
    paragraph = reference_anchor.insert_paragraph_before(text)
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.space_after = Pt(6)

doc.save(path)
