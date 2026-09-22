"""Atkuria Word dokumentą iš bendro plano turinio.

Paleisti: python scripts/build_colloquium_docx.py
Reikia python-docx, lxml, pandoc ir build_colloquium_pdf.py priklausomybių.
Pirmiausia atkuriamas Markdown ir atsarginis PDF, tada Word dokumentas.
Galutinį PDF eksportuoti iš šio DOCX: taip abiejų failų maketas sutampa.
Formulės konvertuojamos į redaguojamus Word OMML objektus, ne paveikslėlius.
"""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import runpy
import subprocess
from zipfile import ZipFile

import fitz
from lxml import etree
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
content = runpy.run_path(str(ROOT / 'scripts/build_colloquium_pdf.py'))
blocks = [block for section in content['pages'] for block in section]
formulas = [value for kind, value in blocks if kind == 'eq']
math_markdown = '\n\n'.join('$$\n' + formula + '\n$$' for formula in formulas)
converted = subprocess.run(
    ['pandoc', '-f', 'markdown', '-t', 'docx'],
    input=math_markdown.encode(), capture_output=True, check=True,
)
with ZipFile(BytesIO(converted.stdout)) as archive:
    math_xml = etree.fromstring(archive.read('word/document.xml'))
math_nodes = math_xml.findall('.//' + qn('m:oMathPara'))
assert len(math_nodes) == 15

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Mm(210), Mm(297)
section.left_margin, section.right_margin = Mm(30), Mm(20)
section.top_margin = section.bottom_margin = Mm(20)
section.footer_distance = Mm(10)
for name in ('Normal', 'Heading 1', 'Heading 2', 'Caption'):
    style = doc.styles[name]
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().rFonts
    for key in list(fonts.attrib):
        if key.endswith('Theme'):
            del fonts.attrib[key]
    for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn('w:' + key), 'Times New Roman')
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_after = Pt(6)
    if name.startswith('Heading'):
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12)
doc.styles['Normal'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
doc.styles['Heading 1'].font.size = Pt(14)
doc.styles['Caption'].font.italic = False
doc.core_properties.title = 'Elektroninės parduotuvės pirkimo ketinimo tyrimo įgyvendinimo planas'
doc.core_properties.author = 'Andrej Kondratjev'


def rich(paragraph, text):
    # lxml XML parseris išlaiko mūsų paprastų žymių poras, taip pat <link> tekstą.
    root = etree.fromstring(('<span>' + text.replace('&nbsp;', '&#160;') + '</span>').encode())

    def walk(node, bold=False, italic=False, sub=False, link=None):
        bold, italic = bold or node.tag == 'b', italic or node.tag == 'i'
        sub = sub or node.tag == 'sub'
        link = node.get('href') if node.tag == 'link' else link

        def add(value):
            if value:
                run = paragraph.add_run(value.replace('`', ''))
                run.bold, run.italic, run.font.subscript = bold, italic, sub
                if link:
                    relation = paragraph.part.relate_to(link, RT.HYPERLINK, is_external=True)
                    hyperlink = OxmlElement('w:hyperlink')
                    hyperlink.set(qn('r:id'), relation)
                    run.font.underline = True
                    hyperlink.append(run._r)
                    paragraph._p.append(hyperlink)
        add(node.text)
        for child in node:
            if child.tag == 'br':
                paragraph.add_run().add_break()
            else:
                walk(child, bold, italic, sub, link)
            add(child.tail)
    walk(root)


def table_layout(table, widths, borders=True):
    table.autofit = False
    for col, width in zip(table.columns, widths):
        col.width = Mm(width)
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for cell, width in zip(row.cells, widths):
            cell.width = Mm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.paragraph_format.space_after = Pt(5)
                paragraph.paragraph_format.space_before = Pt(5)
                for run in paragraph.runs:
                    run.font.size = Pt(10)
    border = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        element = OxmlElement('w:' + edge)
        element.set(qn('w:val'), 'single' if borders else 'nil')
        element.set(qn('w:sz'), '4')
        element.set(qn('w:color'), '999999')
        border.append(element)
    table._tbl.tblPr.append(border)


# Schéma atvaizduojama kaip tikras paveikslas, ne Mermaid programos tekstas.
tmp = ROOT / 'tmp' / 'colloquium'
tmp.mkdir(parents=True, exist_ok=True)
diagram_pdf = tmp / 'flow.pdf'
c = content['canvas'].Canvas(str(diagram_pdf), pagesize=(505, 100))
flow = content['FlowDiagram']('Sprendimo eiga')
flow.drawOn(c, 5, 0)
c.save()
with fitz.open(diagram_pdf) as source:
    source[0].get_pixmap(matrix=fitz.Matrix(3, 3)).save(str(tmp / 'flow.png'))

eq_number = 0
table_number = 0
for kind, value in blocks:
    if kind in ('title', 'h', 'p', 'small'):
        paragraph = doc.add_paragraph(style={'title': 'Heading 1', 'h': 'Heading 2'}.get(kind, 'Normal'))
        rich(paragraph, value)
        if kind == 'small':
            paragraph.paragraph_format.line_spacing = 1.15
            paragraph.paragraph_format.keep_together = True
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if '<br/>' in value:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    elif kind == 'table':
        headers, rows, widths = value
        table_number += 1
        caption = doc.add_paragraph(f'{table_number} lentelė. ' + {
            1: 'Pagrindiniai neaiškumai ir numatyti sprendimai',
            2: 'Projekto etapai ir jų rezultatai',
            3: 'Atsitiktinio miško ir alternatyvų palyginimas',
            4: 'Programos veiksmų seka',
            5: 'Formulių ir programos dalių ryšys',
            6: 'Papildomi bandymai',
            7: 'Grėsmės išvadų galiojimui',
            8: 'AI pasiūlymų ir rezultatų patikra',
        }[table_number], style='Caption')
        caption.paragraph_format.keep_with_next = True
        table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
        for row_i, row_data in enumerate([headers] + rows):
            for cell, text in zip(table.rows[row_i].cells, row_data):
                rich(cell.paragraphs[0], text)
                if row_i == 0:
                    for run in cell.paragraphs[0].runs:
                        run.bold = True
        table_layout(table, [160 * width / sum(widths) for width in widths])
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
    elif kind == 'eq':
        math = deepcopy(math_nodes[eq_number])
        eq_number += 1
        # Pandoc „cases“ kartais palieka numatytuosius skliaustus: aiškiai
        # užrašome kairį riestinį skliaustą ir jokio dešinio skliausto.
        if '\\begin{cases}' in value:
            delimiter = math.find('.//' + qn('m:d'))
            if delimiter is not None:
                props = delimiter.find(qn('m:dPr'))
                if props is None:
                    props = OxmlElement('m:dPr')
                    delimiter.insert(0, props)
                for name, char in [('begChr', '{'), ('endChr', '')]:
                    el = props.find(qn('m:' + name))
                    if el is None:
                        el = OxmlElement('m:' + name)
                        props.append(el)
                    el.set(qn('m:val'), char)
        # Eksportuotojas neteisingai atvaizduoja m:d dešinį skliaustą.
        # Aiškūs redaguojami matematiniai simboliai išlaiko teisingus
        # skliaustus ir Word, ir PDF; dalinės taisyklės lieka OMML matrica.
        for delimiter in reversed(math.findall('.//' + qn('m:d'))):
            props = delimiter.find(qn('m:dPr'))
            def character(name, default):
                item = props.find(qn('m:' + name)) if props is not None else None
                return item.get(qn('m:val'), default) if item is not None else default
            opening, closing = character('begChr', '('), character('endChr', ')')
            replacements = []
            def bracket(text):
                run = OxmlElement('m:r')
                if text == '{':
                    formatting = OxmlElement('w:rPr')
                    size = OxmlElement('w:sz')
                    size.set(qn('w:val'), '56')
                    formatting.append(size)
                    run.append(formatting)
                symbol = OxmlElement('m:t')
                symbol.text = text
                run.append(symbol)
                return run
            if opening:
                replacements.append(bracket(opening))
            for expression in delimiter.findall(qn('m:e')):
                replacements.extend(list(expression))
            if closing:
                replacements.append(bracket(closing))
            parent, index = delimiter.getparent(), delimiter.getparent().index(delimiter)
            parent.remove(delimiter)
            for offset, replacement in enumerate(replacements):
                parent.insert(index + offset, replacement)
        table = doc.add_table(rows=1, cols=2)
        table_layout(table, [145, 15], borders=False)
        left, right = table.rows[0].cells
        left.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        left.paragraphs[0]._p.append(math)
        right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        right.paragraphs[0].add_run(f'({eq_number})')
        for cell in table.rows[0].cells:
            cell.paragraphs[0].paragraph_format.keep_with_next = True
    elif kind == 'diagram':
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(tmp / 'flow.png'), width=Mm(160))

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
output = ROOT / 'docs/KOLIOKVIUMO_PLANAS.docx'
doc.save(output)
print(f'{output}: {eq_number} OMML formulių, {table_number} lentelių')
